from pathlib import Path
from datetime import datetime, timezone
from collections import Counter
import csv, hashlib, importlib.metadata, io, json, os, re, sys, uuid
from urllib.parse import quote
import openpyxl
from pypdf import PdfReader

# Set PROJECT_ROOT when Jupyter starts outside this extracted folder.
PROJECT_ROOT = Path.cwd().resolve()
INPUT = PROJECT_ROOT / "input"
NOTEBOOK = PROJECT_ROOT / "metadata_provenance_atlas.ipynb"
if not INPUT.is_dir() or not NOTEBOOK.is_file():
    raise FileNotFoundError("Set PROJECT_ROOT to the folder containing this notebook and input/.")

# Use a role or a researcher identifier when known. Unknown is not an inferred identity.
CURATOR_ID = "unknown"
PROJECT_ID = "queerdatagap"
CSV_ENCODING = "utf-8-sig"
CSV_DELIMITER = None  # None uses csv.Sniffer over a sample; set explicitly if needed.
HEADER_ROWS = {"QueerDataGap_Lexico_Situado.xlsx": 4, "lista_podcasts.xlsx": 3}
RUN_ID = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S") + "_" + uuid.uuid4().hex[:8]
RUN = PROJECT_ROOT / "runs" / RUN_ID
RUN.mkdir(parents=True, exist_ok=False)

def now():
    return datetime.now(timezone.utc).isoformat()

def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()

def dump(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

def relative(path):
    return Path(path).relative_to(PROJECT_ROOT).as_posix()

def show(rows, limit=10):
    # Plain JSON is also readable outside an interactive notebook.
    print(json.dumps(rows[:limit], ensure_ascii=False, indent=2, default=str))

versions = {name: importlib.metadata.version(name) for name in ["openpyxl", "pypdf"]}
config = {"project_id": PROJECT_ID, "curator_id": CURATOR_ID,
          "csv_encoding": CSV_ENCODING, "csv_delimiter": CSV_DELIMITER,
          "header_rows": HEADER_ROWS, "python": sys.version, "packages": versions}
dump(RUN / "run_config.json", config)
print("Run:", RUN_ID)


manifest = json.loads((PROJECT_ROOT / "input_manifest.json").read_text(encoding="utf-8"))
repo_snapshot = json.loads((PROJECT_ROOT / "repository_snapshot.json").read_text(encoding="utf-8"))
source_lookup = {x["relative_path"]: x for x in manifest + repo_snapshot}
package_checks = []
for entry in manifest + repo_snapshot:
    path = PROJECT_ROOT / entry["relative_path"]
    package_checks.append({"path": entry["relative_path"], "exists": path.is_file(),
                           "hash_matches": path.is_file() and sha256(path) == entry["sha256"]})
if not all(x["hash_matches"] for x in package_checks):
    raise ValueError("An input is missing or changed. Review the manifests before proceeding.")

annotations = {}
catalogue_path = PROJECT_ROOT / "data_catalogue.xlsx"
if catalogue_path.exists():
    book = openpyxl.load_workbook(catalogue_path, read_only=True, data_only=False)
    sheet = book["Data catalogue"]
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    for values in iterator:
        row = dict(zip(headers, values))
        key = row.get("Relative path")
        if not key or not str(key).startswith("input/"):
            continue
        if key in annotations:
            raise ValueError(f"Duplicate catalogue path: {key}")
        annotations[key] = {k: row.get(k) for k in ["Source or origin", "Who creates it",
                           "Who can access it", "Who can modify it", "Evidence or notes"]}
    book.close()
dump(RUN / "curator_annotations.json", annotations)
show(package_checks)


def read_csv_table(path):
    text = Path(path).read_text(encoding=CSV_ENCODING)  # Fail visibly on an incorrect encoding.
    delimiter = CSV_DELIMITER
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not rows:
        return {"row_count": 0, "columns": [], "delimiter": delimiter, "empty_file": True}
    header, body = rows[0], rows[1:]
    return {"row_count": len(body), "column_count": len(header), "columns": header,
            "encoding": CSV_ENCODING, "delimiter": delimiter,
            "duplicate_headers": [k for k, n in Counter(header).items() if n > 1],
            "inconsistent_row_numbers": [i for i, r in enumerate(body, 2) if len(r) != len(header)],
            "empty_cells": sum(v == "" for r in body for v in r)}

def xlsx_tables(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    tables = []
    try:
        for sheet in book:
            rows = list(sheet.iter_rows(values_only=True))
            header_row = HEADER_ROWS.get(Path(path).name, 1)
            header_values = rows[header_row - 1] if len(rows) >= header_row else ()
            # Include columns used anywhere, even when their header is empty.
            width = max((i + 1 for row in rows for i, v in enumerate(row) if v is not None), default=0)
            headers = [str(header_values[i]) if i < len(header_values) and header_values[i] is not None
                       else f"unnamed_column_{i + 1}" for i in range(width)]
            seen = Counter()
            unique = []
            for h in headers:
                seen[h] += 1
                unique.append(h if seen[h] == 1 else f"{h}__{seen[h]}")
            body = [(i, list(row[:width]) + [None] * max(0, width - len(row)))
                    for i, row in enumerate(rows[header_row:], header_row + 1)
                    if any(v is not None for v in row)]
            tables.append({"sheet": sheet.title, "header_row": header_row, "columns": unique,
                           "original_headers": headers, "rows": body,
                           "formula_count": sum(isinstance(v, str) and v.startswith("=")
                                                for _, row in body for v in row),
                           "excluded_blank_rows": len(rows[header_row:]) - len(body)})
    finally:
        book.close()
    return tables

STAMP = r"(\d{2,}):([0-5]\d):([0-5]\d)[,.](\d{3})"
TIMING = re.compile(r"^" + STAMP + r"\s*-->\s*" + STAMP + r"(?:\s+.*)?$")

def parse_srt(path):
    text = Path(path).read_text(encoding="utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
    cues, problems = [], []
    for index, block in enumerate(re.split(r"\n\s*\n", text.strip()), 1):
        if not block.strip():
            continue
        lines = block.splitlines()
        line_index = next((i for i, line in enumerate(lines[:2]) if "-->" in line), None)
        match = TIMING.match(lines[line_index].strip()) if line_index is not None else None
        if not match:
            problems.append({"block": index, "issue": "Malformed or missing timestamp"})
            continue
        parts = list(map(int, match.groups()))
        start = parts[0]*3600 + parts[1]*60 + parts[2] + parts[3]/1000
        end = parts[4]*3600 + parts[5]*60 + parts[6] + parts[7]/1000
        if end < start:
            problems.append({"block": index, "issue": "End precedes start"})
            continue
        cue = {"source_block": index, "cue_id": lines[0].strip() if line_index == 1 else str(index),
               "start_seconds": start, "end_seconds": end,
               "text": "\n".join(lines[line_index + 1:])}
        if cues and start < cues[-1]["end_seconds"]:
            problems.append({"block": index, "issue": "Overlaps or precedes previous cue"})
        if not cue["text"].strip():
            problems.append({"block": index, "issue": "Empty cue text"})
        cues.append(cue)
    return cues, problems

def profile(path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_table(path)
    if suffix == ".srt":
        cues, problems = parse_srt(path)
        return {"cue_count": len(cues), "max_cue_end_seconds": max((r["end_seconds"] for r in cues), default=None),
                "timestamp_issues": problems, "encoding": "utf-8-sig",
                "register_matches_unverified": REGISTER_ENTRIES.get(path.stem, [])}
    if suffix == ".xlsx":
        return {"sheets": [{k: v for k, v in t.items() if k != "rows"} | {"row_count": len(t["rows"])}
                           for t in xlsx_tables(path)]}
    if suffix == ".pdf":
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ValueError("Encrypted PDF requires a password")
        return {"page_count": len(reader.pages), "encrypted": reader.is_encrypted,
                "embedded_metadata_unverified": {str(k): str(v) for k, v in (reader.metadata or {}).items()}}
    if suffix == ".ipynb":
        notebook = json.loads(path.read_text(encoding="utf-8"))
        cells = notebook.get("cells", [])
        code = "\n\n".join("".join(c.get("source", [])) for c in cells if c.get("cell_type") == "code")
        return {"cell_counts": dict(Counter(c.get("cell_type") for c in cells)),
                "kernel": notebook.get("metadata", {}).get("kernelspec", {}),
                "code_source_sha256": hashlib.sha256(code.encode()).hexdigest(), "executed_here": False}
    return {"profile_note": "Generic file metadata only"}


REGISTER_ENTRIES = {}
register_path = INPUT / "tables" / "lista_podcasts.xlsx"
if register_path.exists():
    for table in xlsx_tables(register_path):
        for row_number, values in table["rows"]:
            entry = dict(zip(table["columns"], values))
            identifier = entry.get("Nombre archivo transcripción")
            if identifier:
                key = Path(str(identifier).strip()).stem
                record = {"episode_title": entry.get("Episodio concreto"), "listed_url": entry.get("Enlace"),
                          "listed_platform": entry.get("Plataforma"), "listed_country": entry.get("País"),
                          "assigned_analyst": entry.get("Quién lo analiza"),
                          "evidence": relative(register_path) + f"#sheet={table['sheet']}&row={row_number}"}
                REGISTER_ENTRIES.setdefault(key, []).append(record)

assets, activities = [], []
FORMAT = {".srt": "SRT transcript", ".pdf": "PDF document", ".xlsx": "Excel workbook",
          ".csv": "CSV dataset", ".ipynb": "Jupyter notebook", ".py": "Python code", ".json": "JSON metadata"}

def asset(path, stage="input", origin=None):
    rel, digest = relative(path), sha256(path)
    source = source_lookup.get(rel, {})
    curated = annotations.get(rel, {})
    started = now()
    try:
        technical, status, error = profile(path), "ok", None
    except Exception as exc:
        technical, status, error = {}, "error", f"{type(exc).__name__}: {exc}"
    return {"qualified_name": f"{PROJECT_ID}://artifact/{quote(rel, safe='/')}@sha256:{digest}",
            "name": path.name, "relative_path": rel, "sha256": digest, "size_bytes": path.stat().st_size,
            "format": FORMAT.get(path.suffix.lower(), path.suffix.lstrip('.').upper()), "stage": stage,
            "source_origin": origin or curated.get("Source or origin") or source.get("source_origin", "Unknown"),
            "creators": curated.get("Who creates it") or ("qdg-metadata-notebook-1.0" if stage != "input" else "Unknown"),
            "access": curated.get("Who can access it") or "Unknown",
            "modifiers": curated.get("Who can modify it") or "Unknown",
            "curation_evidence": curated.get("Evidence or notes") or "Needs curator review",
            "filesystem_mtime_utc": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
            "extraction_started_utc": started, "extracted_at_utc": now(), "status": status, "error": error,
            "technical": technical, "git_commit": source.get("commit"),
            "git_blob_sha1": source.get("git_blob_sha1")}

for path in sorted(INPUT.rglob("*")):
    if path.is_file() and not path.name.startswith("."):
        assets.append(asset(path))

# Preserve the executable source used in this run, independent of notebook output cells.
notebook_json = json.loads(NOTEBOOK.read_text(encoding="utf-8"))
code_source = "\n\n".join("".join(c["source"]) for c in notebook_json["cells"] if c["cell_type"] == "code")
code_path = RUN / "executed_notebook_source.py"
code_path.write_text(code_source, encoding="utf-8")
code_asset = asset(code_path, "run specification", "Code cells of metadata_provenance_atlas.ipynb")
config_asset = asset(RUN / "run_config.json", "run specification", "Configuration recorded by this execution")
annotations_asset = asset(RUN / "curator_annotations.json", "run specification", "Curator annotations loaded for this run")
assets.extend([code_asset, config_asset, annotations_asset])
# The manifests supply origin and version statements used in the metadata output.
for name in ["input_manifest.json", "repository_snapshot.json"]:
    manifest_asset = asset(PROJECT_ROOT / name, "run specification", "Recorded during package assembly")
    manifest_asset["creators"] = "Package assembly process"
    assets.append(manifest_asset)
show([{k: a[k] for k in ["name", "format", "size_bytes", "status"]} for a in assets])


def write_csv(path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)

def record_process(name, inputs, outputs, started, parameters):
    event = {"id": f"{PROJECT_ID}://run/{RUN_ID}/process/{len(activities) + 1}",
             "name": name, "started_at_utc": started, "ended_at_utc": now(),
             "inputs": [a["qualified_name"] for a in inputs],
             "outputs": [a["qualified_name"] for a in outputs],
             "agent": "qdg-metadata-notebook-1.0", "curator_id": CURATOR_ID,
             "parameters": parameters, "code_sha256": code_asset["sha256"],
             "status": "completed"}
    activities.append(event)

for original in list(assets):
    path = PROJECT_ROOT / original["relative_path"]
    if original["stage"] != "input" or original["status"] != "ok":
        continue
    started = now()
    if path.suffix.lower() == ".xlsx":
        for index, t in enumerate(xlsx_tables(path), 1):
            started = now()
            # Numbered names avoid collisions caused by accented or punctuation-heavy sheet names.
            destination = RUN / f"{path.stem}__sheet_{index}.csv"
            write_csv(destination, ["_source_row"] + t["columns"], [[rownum] + row for rownum, row in t["rows"]])
            derived = asset(destination, "derived", original["qualified_name"] + "#sheet=" + quote(t["sheet"]))
            assets.append(derived)
            record_process("Export spreadsheet sheet to CSV", [original, code_asset, config_asset], [derived], started,
                           {"sheet": t["sheet"], "header_row": t["header_row"], "blank_rows": "excluded",
                            "formulas": "preserved as expressions, not recalculated", "encoding": "utf-8-sig"})
    elif path.suffix.lower() == ".srt":
        cues, issues = parse_srt(path)
        destination = RUN / f"{path.stem}__cues.csv"
        fields = ["source_block", "cue_id", "start_seconds", "end_seconds", "text"]
        write_csv(destination, fields, [[cue[k] for k in fields] for cue in cues])
        derived = asset(destination, "derived", original["qualified_name"])
        assets.append(derived)
        record_process("Parse SRT cues to CSV", [original, code_asset, config_asset], [derived], started,
                       {"timestamp_unit": "seconds", "text": "preserved", "reported_issues": issues})

# Create a metadata artifact generated by this extraction run. It does not contain itself.
started = now()
metadata_path = RUN / "extracted_metadata.json"
dump(metadata_path, assets)
metadata_asset = asset(metadata_path, "metadata output", "Metadata extracted in run " + RUN_ID)
record_process("Extract and assemble metadata", list(assets), [metadata_asset],
               min(a["extraction_started_utc"] for a in assets), {"versions": versions, "profilers": "format-specific"})
assets.append(metadata_asset)
dump(RUN / "asset_inventory.json", assets)
dump(RUN / "process_events.json", activities)
print(f"Recorded {len(assets)} artifacts and {len(activities)} completed processes.")


prov = {"prefix": {"prov": "http://www.w3.org/ns/prov#", "qdg": "https://example.org/queerdatagap/"},
        "entity": {}, "activity": {}, "agent": {}, "used": {}, "wasGeneratedBy": {},
        "wasAssociatedWith": {}, "wasDerivedFrom": {}}
entity_ids = {a["qualified_name"]: f"qdg:e{i}" for i, a in enumerate(assets, 1)}
for a in assets:
    prov["entity"][entity_ids[a["qualified_name"]]] = {
        "prov:label": a["name"], "qdg:qualifiedName": a["qualified_name"], "qdg:sha256": a["sha256"]}
prov["agent"]["qdg:software"] = {"prov:type": {"$": "prov:SoftwareAgent", "type": "prov:QUALIFIED_NAME"},
                                          "prov:label": "Metadata notebook 1.0"}
if CURATOR_ID != "unknown":
    prov["agent"]["qdg:curator"] = {"prov:label": CURATOR_ID}
for i, event in enumerate(activities, 1):
    aid = f"qdg:a{i}"
    prov["activity"][aid] = {"prov:label": event["name"], "prov:startTime": event["started_at_utc"],
                            "prov:endTime": event["ended_at_utc"], "qdg:parameters": json.dumps(event["parameters"])}
    prov["wasAssociatedWith"][f"qdg:assoc{i}"] = {"prov:activity": aid, "prov:agent": "qdg:software"}
    if CURATOR_ID != "unknown":
        prov["wasAssociatedWith"][f"qdg:curator{i}"] = {"prov:activity": aid, "prov:agent": "qdg:curator"}
    for j, name in enumerate(event["inputs"]):
        prov["used"][f"qdg:u{i}_{j}"] = {"prov:activity": aid, "prov:entity": entity_ids[name]}
    for k, name in enumerate(event["outputs"]):
        prov["wasGeneratedBy"][f"qdg:g{i}_{k}"] = {"prov:entity": entity_ids[name], "prov:activity": aid}
        for j, source in enumerate(event["inputs"]):
            if source in [code_asset["qualified_name"], config_asset["qualified_name"]]:
                continue
            prov["wasDerivedFrom"][f"qdg:d{i}_{k}_{j}"] = {
                "prov:generatedEntity": entity_ids[name], "prov:usedEntity": entity_ids[source], "prov:activity": aid}
dump(RUN / "provenance.prov.json", prov)
print("PROV entities:", len(prov["entity"]), "activities:", len(prov["activity"]))


ARTIFACT_TYPE = "qdg_artifact_v1"
PROCESS_TYPE = "qdg_process_v1"

def attribute(name, kind="string", optional=True):
    return {"name": name, "typeName": kind, "isOptional": optional, "cardinality": "SINGLE",
            "valuesMinCount": 0 if optional else 1, "valuesMaxCount": 1,
            "isUnique": False, "isIndexable": False}

artifact_fields = ["relativePath", "sha256", "fileFormat", "stage", "sourceOrigin", "creators",
                   "accessRights", "modificationRights", "curationEvidence", "filesystemMtimeUtc",
                   "extractedAtUtc", "extractionStatus", "extractionError", "technicalMetadataJson",
                   "gitCommit", "gitBlobSha1"]
process_fields = ["runId", "startedAtUtc", "endedAtUtc", "softwareAgent", "curatorId",
                  "parametersJson", "codeSha256", "executionStatus"]
typedefs = {"enumDefs": [], "structDefs": [], "classificationDefs": [], "relationshipDefs": [],
           "entityDefs": [
               {"name": ARTIFACT_TYPE, "superTypes": ["DataSet"], "typeVersion": "1.0",
                "description": "A versioned research artifact in the curation workflow",
                "attributeDefs": [attribute(x) for x in artifact_fields] + [attribute("sizeBytes", "long")]},
               {"name": PROCESS_TYPE, "superTypes": ["Process"], "typeVersion": "1.0",
                "description": "An observed execution with input and output lineage",
                "attributeDefs": [attribute(x) for x in process_fields]}]}
guids = {a["qualified_name"]: str(-i) for i, a in enumerate(assets, 1)}
entities = []
mapping = {"relative_path": "relativePath", "sha256": "sha256", "size_bytes": "sizeBytes",
           "format": "fileFormat", "stage": "stage", "source_origin": "sourceOrigin", "creators": "creators",
           "access": "accessRights", "modifiers": "modificationRights", "curation_evidence": "curationEvidence",
           "filesystem_mtime_utc": "filesystemMtimeUtc", "extracted_at_utc": "extractedAtUtc",
           "status": "extractionStatus", "error": "extractionError", "git_commit": "gitCommit", "git_blob_sha1": "gitBlobSha1"}
for a in assets:
    attrs = {target: a[source] for source, target in mapping.items() if a.get(source) is not None}
    attrs.update({"qualifiedName": a["qualified_name"], "name": a["name"],
                  "technicalMetadataJson": json.dumps(a["technical"], ensure_ascii=False)})
    entities.append({"typeName": ARTIFACT_TYPE, "guid": guids[a["qualified_name"]], "attributes": attrs})
for i, event in enumerate(activities, len(assets) + 1):
    attrs = {"qualifiedName": event["id"], "name": event["name"], "runId": RUN_ID,
             "startedAtUtc": event["started_at_utc"], "endedAtUtc": event["ended_at_utc"],
             "softwareAgent": event["agent"], "curatorId": CURATOR_ID,
             "parametersJson": json.dumps(event["parameters"], ensure_ascii=False),
             "codeSha256": event["code_sha256"], "executionStatus": event["status"]}
    for direction in ["inputs", "outputs"]:
        attrs[direction] = [{"guid": guids[name], "typeName": ARTIFACT_TYPE} for name in event[direction]]
    entities.append({"typeName": PROCESS_TYPE, "guid": str(-i), "attributes": attrs})
payload = {"entities": entities}
dump(RUN / "atlas_typedefs.json", typedefs)
dump(RUN / "atlas_entities.json", payload)
print("Atlas payload:", len(entities), "entities")


qnames = [e["attributes"]["qualifiedName"] for e in entities]
assert len(qnames) == len(set(qnames)), "Duplicate entity qualified names"
assert all(e["attributes"].get("name") for e in entities)
assert all(re.fullmatch(r"[0-9a-f]{64}", a["sha256"]) for a in assets)
for event in activities:
    assert set(event["inputs"] + event["outputs"]) <= set(guids)
    assert not set(event["inputs"]) & set(event["outputs"])
    assert datetime.fromisoformat(event["started_at_utc"]) <= datetime.fromisoformat(event["ended_at_utc"])
for entry in manifest + repo_snapshot:
    assert sha256(PROJECT_ROOT / entry["relative_path"]) == entry["sha256"], "Input changed during run"
review = {"artifact_count": len(assets), "process_count": len(activities),
          "extraction_errors": [{"path": a["relative_path"], "error": a["error"]} for a in assets if a["status"] == "error"],
          "srt_issues": [{"path": a["relative_path"], "issues": a["technical"]["timestamp_issues"]}
                         for a in assets if a["technical"].get("timestamp_issues")],
          "unknown_creators": sum(a["creators"].lower() == "unknown" for a in assets),
          "local_reference_checks": "passed", "atlas_server_validation": "not performed"}
dump(RUN / "validation_report.json", review)
write_csv(RUN / "data_catalogue_snapshot.csv",
          ["qualified_name", "name", "source_origin", "format", "creators", "access", "modifiers", "relative_path", "sha256", "status"],
          [[a[k] for k in ["qualified_name", "name", "source_origin", "format", "creators", "access", "modifiers", "relative_path", "sha256", "status"]]
          for a in assets])
show([review])


SUBMIT_TO_ATLAS = False

def submit_to_atlas():
    import base64, urllib.request, urllib.error
    base = os.environ.get("ATLAS_URL", "").rstrip("/")
    user = os.environ.get("ATLAS_USER", "")
    password = os.environ.get("ATLAS_PASSWORD", "")
    if not base.startswith("https://") or not user or not password:
        raise ValueError("Set HTTPS ATLAS_URL, ATLAS_USER and ATLAS_PASSWORD before submission.")
    auth = base64.b64encode(f"{user}:{password}".encode()).decode()

    def request(method, path, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(base + "/api/atlas/v2" + path, data=data, method=method,
                                     headers={"Authorization": "Basic " + auth,
                                              "Content-Type": "application/json", "Accept": "application/json"})
        with urllib.request.urlopen(req, timeout=60) as response:
            return json.load(response)

    missing = []
    for desired in typedefs["entityDefs"]:
        try:
            current = request("GET", "/types/entitydef/name/" + desired["name"])
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                missing.append(desired)
                continue
            raise
        actual = {a["name"]: a["typeName"] for a in current.get("attributeDefs", [])}
        if not set(desired["superTypes"]) <= set(current.get("superTypes", [])) or any(
            actual.get(a["name"]) != a["typeName"] for a in desired["attributeDefs"]
        ):
            raise ValueError(f"Atlas type mismatch: {desired['name']}. Review a versioned type migration.")
    if missing:
        request("POST", "/types/typedefs", {"entityDefs": missing})
    response = request("POST", "/entity/bulk", payload)
    dump(RUN / "atlas_submission_response.json", response)
    print("Atlas submission completed; response saved locally.")

if SUBMIT_TO_ATLAS:
    submit_to_atlas()
else:
    print("Local export complete. Atlas submission is disabled.")
